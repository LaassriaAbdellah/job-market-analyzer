
import io
import json
import re
import unicodedata
from pathlib import Path

import fitz
from docx import Document
from fastapi import HTTPException


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

KEYWORDS_PATH = BASE_DIR / "config" / "keywords.json"
ALIASES_PATH = BASE_DIR / "config" / "skill_aliases.json"

MAX_FILE_SIZE = 5 * 1024 * 1024


# ============================================================
# EXTRACTION DU TEXTE PDF
# ============================================================

def extract_pdf_text(content: bytes) -> str:
    """
    Extrait le texte d'un CV au format PDF.
    """

    try:
        with fitz.open(
            stream=content,
            filetype="pdf"
        ) as document:

            texts = []

            for page in document:
                page_text = page.get_text("text")

                if page_text:
                    texts.append(page_text)

            return "\n".join(texts)

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Impossible de lire ce PDF."
        ) from exc


# ============================================================
# EXTRACTION DU TEXTE DOCX
# ============================================================

def extract_docx_text(content: bytes) -> str:
    """
    Extrait le texte des paragraphes et tableaux
    d'un document Word.
    """

    try:
        document = Document(io.BytesIO(content))

        parts = []

        # Paragraphes
        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Tableaux
        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    if cell.text.strip():
                        parts.append(cell.text)

        return "\n".join(parts)

    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Impossible de lire ce DOCX."
        ) from exc


# ============================================================
# VALIDATION ET EXTRACTION DU CV
# ============================================================

def extract_cv_text(
    filename: str,
    content: bytes
) -> str:
    """
    Vérifie le format et la taille du CV,
    puis extrait son contenu textuel.
    """

    # Vérifier la taille
    if len(content) > MAX_FILE_SIZE:

        raise HTTPException(
            status_code=413,
            detail="Le CV ne doit pas dépasser 5 Mo."
        )

    # Vérifier le format
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":

        text = extract_pdf_text(content)

    elif extension == ".docx":

        text = extract_docx_text(content)

    else:

        raise HTTPException(
            status_code=415,
            detail="Formats acceptés : PDF et DOCX."
        )

    # Vérifier la présence de texte
    if not text.strip():

        raise HTTPException(
            status_code=422,
            detail=(
                "Aucun texte détecté dans le CV. "
                "Le document est peut-être scanné "
                "ou constitué uniquement d'images."
            )
        )

    return text


# ============================================================
# NORMALISATION DU TEXTE
# ============================================================

def normalize_text(text: str) -> str:
    """
    Normalise le texte pour faciliter
    la comparaison des compétences.

    - Conversion en minuscules
    - Suppression des accents
    - Uniformisation des espaces
    """

    text = unicodedata.normalize(
        "NFKD",
        text
    )

    text = "".join(
        char
        for char in text
        if not unicodedata.combining(char)
    )

    text = text.casefold()

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ============================================================
# CHARGEMENT DES DICTIONNAIRES
# ============================================================

def load_skill_dictionaries() -> tuple[dict, dict]:
    """
    Charge les compétences et leurs synonymes.
    """

    try:
        with KEYWORDS_PATH.open(
            "r",
            encoding="utf-8"
        ) as file:

            keywords = json.load(file)

        with ALIASES_PATH.open(
            "r",
            encoding="utf-8"
        ) as file:

            aliases = json.load(file)

    except FileNotFoundError as exc:

        raise HTTPException(
            status_code=500,
            detail=(
                "Fichier de configuration introuvable : "
                f"{exc.filename}"
            )
        ) from exc

    except json.JSONDecodeError as exc:

        raise HTTPException(
            status_code=500,
            detail="Erreur dans un fichier JSON de configuration."
        ) from exc

    if not isinstance(keywords, dict) or not isinstance(aliases, dict):

        raise HTTPException(
            status_code=500,
            detail="Format invalide des dictionnaires de compétences."
        )

    return keywords, aliases


# ============================================================
# CONSTRUCTION DU DICTIONNAIRE NORMALISÉ
# ============================================================

def build_skill_mapping(
    keywords: dict,
    aliases: dict
) -> dict[str, str]:
    """
    Construit un dictionnaire :

    variante normalisée -> compétence canonique

    Exemple :
        sklearn -> Scikit-learn
        postgres -> PostgreSQL
        scikit learn -> Scikit-learn
    """

    skill_mapping = {}

    # Ajouter les compétences canoniques et synonymes
    for canonical, variants in aliases.items():

        canonical = str(canonical).strip()

        if not canonical:
            continue

        all_variants = [canonical]

        if isinstance(variants, list):
            all_variants.extend(variants)

        for variant in all_variants:

            normalized = normalize_text(str(variant))

            if normalized:
                skill_mapping[normalized] = canonical

    # Ajouter les compétences de keywords.json
    for category, skills in keywords.items():

        if not isinstance(skills, list):
            continue

        for skill in skills:

            skill = str(skill).strip()

            if not skill:
                continue

            normalized = normalize_text(skill)

            # Ne pas écraser une correspondance canonique
            # déjà définie dans skill_aliases.json.
            skill_mapping.setdefault(
                normalized,
                skill
            )

    return skill_mapping


# ============================================================
# EXTRACTION DES COMPÉTENCES
# ============================================================

def extract_cv_skills(text: str) -> list[str]:
    """
    Extrait les compétences techniques du CV.

    Fonctionnalités :
    - Recherche par mots-clés
    - Reconnaissance des synonymes
    - Normalisation des accents et de la casse
    - Suppression des doublons
    """

    keywords, aliases = load_skill_dictionaries()

    skill_mapping = build_skill_mapping(
        keywords,
        aliases
    )

    normalized_cv = normalize_text(text)

    found_skills = {}

    for variant, canonical in skill_mapping.items():

        # Éviter les correspondances à l'intérieur
        # d'autres mots.
        pattern = (
            r"(?<!\w)"
            + re.escape(variant)
            + r"(?!\w)"
        )

        if re.search(pattern, normalized_cv):

            found_skills[canonical.casefold()] = canonical

    return sorted(
        found_skills.values(),
        key=str.casefold
    )
